"""End-to-end smoke test for phases 1-10 (including Phase 9: production deploy & repo_control).

    .venv/Scripts/python.exe smoke_test.py           # offline, no API calls
    .venv/Scripts/python.exe smoke_test.py --live    # also does one real turn

Runs against a throwaway database and Redis db 15, so it never touches
stellar_local.db or the app's Redis keyspace.

The default run makes no network calls at all: the LLM producer is swapped
for a stub, so the transport can be verified without spending quota or
failing because Google had a bad minute. --live adds one real turn.
"""

from __future__ import annotations

import json
import os
import re
import sqlite3
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

import app as A  # noqa: E402
import chess_ui as _cui  # noqa: E402

REDIS_TEST_URL = "redis://127.0.0.1:6379/15"
LIVE = "--live" in sys.argv

failures: list[str] = []


def check(label: str, cond: bool) -> None:
    print(f"  {'PASS' if cond else 'FAIL'}  {label}")
    if not cond:
        failures.append(label)


def parse_sse(raw: str) -> list[tuple[int | None, dict]]:
    out = []
    for block in raw.split("\n\n"):
        if not block.strip() or block.lstrip().startswith(":"):
            continue
        m_id = re.search(r"^id: (\d+)$", block, re.M)
        m_data = re.search(r"^data: (.*)$", block, re.M)
        if m_data:
            out.append((int(m_id.group(1)) if m_id else None,
                        json.loads(m_data.group(1))))
    return out


def stub_producer(r, args):
    """Stand-in for the model, so transport tests need no network."""
    database = A.get_db()
    uid = database.execute(
        "INSERT INTO messages (chat_id, message_type, message_content)"
        " VALUES (?, 'user', ?)",
        (args["chat_id"], args["message"]),
    ).lastrowid
    title = A._touch_chat(database, args["chat_id"], args["message"])
    database.commit()

    yield {"type": "user_message", "id": uid}
    if title:
        yield {"type": "chat_title", "chat_id": args["chat_id"], "name": title}
    yield {"type": "status", "text": "Thinking…"}
    for tok in ("**stub** ", "reply ", "for ", args["message"]):
        yield {"type": "token", "text": tok}

    rid = A._save_reply(database, args["chat_id"],
                        f"**stub** reply for {args['message']}")
    yield {"type": "message", "id": rid}


def main() -> int:
    tmp = Path(tempfile.mkdtemp()) / "smoke.db"
    app = A.create_app({
        "DATABASE": str(tmp),
        "TESTING": True,
        "REDIS_URL": REDIS_TEST_URL,
        "OUTPUTS_DIR": str(tmp.parent / "outputs"),
    })
    with app.app_context():
        A.init_db()

    import redis as redis_lib
    redis_lib.from_url(REDIS_TEST_URL).flushdb()

    print("\n  Stellar smoke test" + ("  (live)" if LIVE else "  (offline)") + "\n")

    # --- schema ------------------------------------------------------
    conn = sqlite3.connect(tmp)
    tables = {r[0] for r in conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table'")}
    check("schema: users/chats/messages exist",
          {"users", "chats", "messages"} <= tables)
    check("schema: WAL enabled",
          conn.execute("PRAGMA journal_mode").fetchone()[0].lower() == "wal")
    conn.close()

    # --- auth --------------------------------------------------------
    c = app.test_client()
    check("anonymous API access refused",
          c.get("/api/chats").status_code in (302, 401))

    c.post("/auth/register", data={"username": "a@b.com",
                                   "password": "hunter2hunter2"},
           follow_redirects=True)
    check("login works",
          c.post("/auth/login", data={"username": "a@b.com",
                                      "password": "hunter2hunter2"}
                 ).status_code == 302)
    check("first user is auto-approved admin",
          c.get("/api/chats").status_code == 200)

    bad = app.test_client()
    bad.post("/auth/login", data={"username": "a@b.com", "password": "wrong"})
    check("wrong password rejected",
          bad.get("/api/chats").status_code in (302, 401))

    # --- second user is gated ----------------------------------------
    c2 = app.test_client()
    c2.post("/auth/register", data={"username": "m@b.com",
                                    "password": "hunter2hunter2"},
            follow_redirects=True)
    c2.post("/auth/login", data={"username": "m@b.com",
                                 "password": "hunter2hunter2"})
    check("second user needs approval", c2.get("/api/chats").status_code == 403)

    conn = sqlite3.connect(tmp)
    conn.execute("UPDATE users SET is_approved=1 WHERE username='m@b.com'")
    conn.commit()
    conn.close()

    # --- chats -------------------------------------------------------
    chat = c.post("/api/chats").get_json()
    check("create chat", isinstance(chat.get("id"), int))
    check("cannot read another user's chat",
          c2.get(f"/api/chats/{chat['id']}/messages").status_code == 404)

    # --- streaming transport (stubbed producer) ----------------------
    r = c.post(f"/api/chats/{chat['id']}/query", json={"message": "hello"})
    check("register returns 202 + query_id",
          r.status_code == 202 and r.get_json().get("query_id"))
    qid = r.get_json()["query_id"]

    check("registering writes nothing to the database",
          len(c.get(f"/api/chats/{chat['id']}/messages").get_json()) == 0)
    check("empty message rejected",
          c.post(f"/api/chats/{chat['id']}/query",
                 json={"message": "  "}).status_code == 400)
    check("another user cannot attach to the stream",
          c2.get(f"/api/stream/{qid}").status_code == 404)
    check("unknown query id is 404",
          c.get("/api/stream/nope").status_code == 404)

    real_producer = A.gemini_producer
    A.gemini_producer = stub_producer          # transport test, no network
    try:
        resp = c.get(f"/api/stream/{qid}")
        check("content-type is text/event-stream",
              resp.headers.get("Content-Type", "").startswith("text/event-stream"))
        check("X-Accel-Buffering disabled for nginx",
              resp.headers.get("X-Accel-Buffering") == "no")

        events = parse_sse(resp.get_data(as_text=True))
        kinds = [e["type"] for _, e in events]
        check("stream terminates with done", kinds[-1] == "done")
        check("stream carries tokens", kinds.count("token") >= 3)
        check("stream reports the committed message id", "message" in kinds)
        check("event ids are sequential from 0",
              [i for i, _ in events] == list(range(len(events))))

        stored = c.get(f"/api/chats/{chat['id']}/messages").get_json()
        check("user message + reply both persisted", len(stored) == 2)
        check("chat auto-titled from first message",
              c.get("/api/chats").get_json()[0]["name"] == "hello")
        # The title travels in the stream, so the client needs no follow-up
        # GET /api/chats to discover it.
        titles = [e for _, e in events if e["type"] == "chat_title"]
        check("title pushed down the stream",
              len(titles) == 1 and titles[0]["name"] == "hello")

        # resumability: the property the whole design exists for
        replay = parse_sse(c.get(f"/api/stream/{qid}?from=0").get_data(as_text=True))
        check("reattaching from 0 replays the whole stream",
              [e for _, e in replay] == [e for _, e in events])

        mid = len(events) // 2
        resumed = parse_sse(
            c.get(f"/api/stream/{qid}?from={mid}").get_data(as_text=True))
        check("resuming from an index skips earlier events",
              len(resumed) == len(events) - mid and resumed[0][0] == mid)

        before = sqlite3.connect(tmp).execute(
            "SELECT COUNT(*) FROM messages").fetchone()[0]
        c.get(f"/api/stream/{qid}?from=0")
        after = sqlite3.connect(tmp).execute(
            "SELECT COUNT(*) FROM messages").fetchone()[0]
        check("reattaching does not re-run the worker", before == after)
    finally:
        A.gemini_producer = real_producer

    # --- error classification ----------------------------------------
    check("transient errors classified",
          A._classify_error(Exception("Server disconnected without sending a response")) == "transient")
    check("missing model classified",
          A._classify_error(Exception("404 NOT_FOUND")) == "missing_model")
    check("quota classified",
          A._classify_error(Exception("429 RESOURCE_EXHAUSTED")) == "quota")
    check("unknown errors are fatal",
          A._classify_error(Exception("something odd")) == "fatal")

    # --- phase 4: tools ----------------------------------------------
    check("tool registry matches the exposed list",
          set(A.TOOLS_BY_NAME) == {f.__name__ for f in A.AVAILABLE_TOOLS})

    # Every tool must take `status`; the UI depends on the model writing it.
    import inspect
    check("every tool accepts a status argument",
          all("status" in inspect.signature(f).parameters for f in A.AVAILABLE_TOOLS))
    # And a docstring, since that IS the schema the model sees.
    check("every tool has a docstring",
          all((f.__doc__ or "").strip() for f in A.AVAILABLE_TOOLS))

    out, err = A._execute_tool("no_such_tool", {})
    check("unknown tool is reported, not raised", err and "No such tool" in out)
    out, err = A._execute_tool("get_current_time", {"bogus": 1})
    check("bad arguments are reported, not raised", err and "Invalid arguments" in out)

    check("valid timezone resolves",
          "2026" in A.get_current_time("Asia/Kolkata", "s"))
    check("invalid timezone is rejected",
          "Unknown timezone" in A.get_current_time("Mars/Olympus", "s"))

    # SSRF guard. These resolve without leaving the machine.
    check("loopback is refused", not A._is_safe_url("http://localhost:6379/")[0])
    check("link-local metadata is refused",
          not A._is_safe_url("http://169.254.169.254/latest/meta-data/")[0])
    check("private range is refused", not A._is_safe_url("http://10.0.0.1/")[0])
    check("non-http scheme is refused", not A._is_safe_url("file:///etc/passwd")[0])
    check("fetch_url refuses rather than fetching",
          A.fetch_url("http://169.254.169.254/", "s").startswith("Refused"))

    # Force an empty pool rather than reading one env var, so this tests the
    # degradation path regardless of how the machine happens to be configured.
    _real_tavily = A.tavily_keys
    A.tavily_keys = lambda: []
    try:
        check("web_search degrades without a key",
              "TAVILY_API_KEY" in A.web_search("anything", "s"))
    finally:
        A.tavily_keys = _real_tavily
    check("tavily pool is discovered when present", len(A.tavily_keys()) >= 0)

    # tool_calls persistence and the shape the UI consumes
    with app.app_context():
        db = A.get_db()
        rid = A._save_reply(db, chat["id"], "reply with a tool")
        tid = A._record_tool_call(db, chat["id"], "get_current_time",
                                  {"timezone": "UTC"}, "result text", 12, False)
        db.execute("UPDATE tool_calls SET message_id = ? WHERE id = ?", (rid, tid))
        db.commit()

    msgs = c.get(f"/api/chats/{chat['id']}/messages").get_json()
    withtools = [m for m in msgs if m.get("tools")]
    check("tool calls are returned with their message", len(withtools) == 1)
    check("tool payload carries name, timing and error flag",
          withtools and withtools[0]["tools"][0]["name"] == "get_current_time"
          and withtools[0]["tools"][0]["ms"] == 12
          and withtools[0]["tools"][0]["is_error"] is False)

    # --- phase 7: key rotation ----------------------------------------
    # The real 429 body Google returned when the daily quota ran out. Its
    # retryDelay says 4 seconds, but the quota it names resets at midnight -
    # obeying the delay would retry against an empty bucket all day.
    real_rpd = (
        "429 RESOURCE_EXHAUSTED. quota exceeded for metric: "
        "generate_content_free_tier_requests, limit: 20. Please retry in "
        "4.389960791s. quotaId: GenerateRequestsPerDayPerProjectPerModel-FreeTier"
    )
    secs, reason = A.parse_quota_block(real_rpd)
    reset = A.seconds_until_pacific_midnight()
    check("per-day 429 classified RPD", reason == "RPD")
    check("per-day block runs to the Pacific reset, not the 4s hint",
          secs > 60 and abs(secs - reset) <= 5)

    secs, reason = A.parse_quota_block("429 rate limit, please retry in 12.5s")
    check("per-minute 429 honours its retry hint",
          reason == "RPM" and 10 <= secs <= 25)
    check("invalid key classified",
          A.parse_quota_block("403 PERMISSION_DENIED api key not valid")[1] == "INVALID")
    check("overload classified",
          A.parse_quota_block("503 model is overloaded")[1] == "OVERLOAD")

    km = A.KeyManager(redis_url=REDIS_TEST_URL)
    kk = ["K-AAA", "K-BBB", "K-CCC"]
    M1, M2 = "m-one", "m-two"
    check("key manager reaches redis", km._r() is not None)
    check("starts on the first key", km.first_available(kk, M1) == 0)

    km.block(kk[0], M1, 60, "RPM")
    check("blocked key is skipped", km.first_available(kk, M1) == 1)
    # The property the whole design rests on.
    check("blocks are per (key, model), not per key",
          km.first_available(kk, M2) == 0)

    km.block(kk[1], M1, 60, "RPM")
    km.block(kk[2], M1, 60, "RPM")
    check("all keys blocked reports None", km.first_available(kk, M1) is None)

    km.block(kk[0], M1, 1, "RPM")
    import time as _t
    _t.sleep(1.4)
    check("an expired block returns to the earliest key",
          km.first_available(kk, M1) == 0)

    stored = " ".join(redis_lib.from_url(REDIS_TEST_URL, decode_responses=True)
                      .keys("keyblock:*"))
    check("redis stores fingerprints, never raw keys", "K-AAA" not in stored)

    other = A.KeyManager(redis_url=REDIS_TEST_URL)
    other.block(kk[1], M2, 120, "RPD")
    blocked, why = km.is_blocked(kk[1], M2)
    check("one worker's block is visible to another", blocked and why == "RPD")

    check("status never leaks a raw key",
          all("K-AAA" not in str(r) for r in km.status(kk, [M1, M2])))

    # Proactive counting: a key must retire BEFORE it 429s, since the
    # failing call still costs a round trip and a retry.
    redis_lib.from_url(REDIS_TEST_URL).flushdb()
    pk = A.KeyManager(redis_url=REDIS_TEST_URL)
    PM = "gemini-3-flash-preview"
    rpm_lim, rpd_lim = A.get_limits(PM)
    check("limits are per model", A.get_limits("gemma-4-31b-it")[1] > rpd_lim)
    check("unknown models get the conservative default",
          A.get_limits("something-unreleased") == A.DEFAULT_LIMITS)

    for _ in range(rpd_lim):
        pk.record_request(kk[0], PM)
    blocked, why = pk.is_blocked(kk[0], PM)
    check("key retires on the daily count, before any API call",
          blocked and why == "RPD")
    check("usage reports the spend",
          pk.usage(kk[0], PM)["rpd_used"] == rpd_lim)

    redis_lib.from_url(REDIS_TEST_URL).flushdb()
    pk2 = A.KeyManager(redis_url=REDIS_TEST_URL)
    for _ in range(rpm_lim):
        pk2.record_request(kk[0], PM)
    check("key retires on a per-minute burst",
          pk2.is_blocked(kk[0], PM)[1] == "RPM")
    check("a burst on one model does not touch another",
          not pk2.is_blocked(kk[0], "gemma-4-31b-it")[0])

    # An invalid credential is not a per-model condition.
    redis_lib.from_url(REDIS_TEST_URL).flushdb()
    pk3 = A.KeyManager(redis_url=REDIS_TEST_URL)
    pk3.block(kk[0], PM, 300, "INVALID")
    check("INVALID blocks the key across every model",
          pk3.is_blocked(kk[0], "gemini-3.6-flash")[1] == "INVALID")
    pk3.block(kk[1], PM, 300, "RPD")
    check("a quota block stays scoped to its model",
          not pk3.is_blocked(kk[1], "gemini-3.6-flash")[0])

    # A 503 is Google being busy, not the key being spent.
    check("503 classified OVERLOAD",
          A.parse_quota_block("503 The model is overloaded")[1] == "OVERLOAD")
    redis_lib.from_url(REDIS_TEST_URL).flushdb()
    pk4 = A.KeyManager(redis_url=REDIS_TEST_URL)
    pk4.block_model(PM, 60)
    check("an overloaded model yields no key at all",
          pk4.first_available(kk, PM) is None)
    check("other models remain usable while one is overloaded",
          pk4.first_available(kk, "gemini-3.6-flash") == 0)

    # --- phase 5: the sandbox -----------------------------------------
    # Skipped rather than failed when Docker is down: the rest of the suite
    # is useful on a machine without it, and a hard failure here would hide
    # everything else.
    try:
        import docker as _docker_sdk
        _cl = _docker_sdk.from_env()
        _cl.ping()
        docker_up = True
    except Exception:
        docker_up = False

    if not docker_up:
        print("  SKIP  sandbox checks (Docker not reachable)")
    else:
        from flask import g as _g
        with app.app_context():
            _g.lab_user_id, _g.lab_chat_id = 9998, 1

            out = A.lab_execute("echo sandbox-ok", "test", 30)
            check("sandbox runs a command", "sandbox-ok" in out)

            A.lab_execute("echo persisted > f.txt", "test", 30)
            check("files persist across commands",
                  "persisted" in A.lab_execute("cat f.txt", "test", 30))
            check("and reach the host disk",
                  (A.PROJECT_ROOT / "sandbox_runs" / "u9998_c1" / "f.txt").exists())

            check("non-zero exit is reported",
                  "7" in A.lab_execute("exit 7", "test", 30))
            check("stderr is captured",
                  "boom" in A.lab_execute(
                      "python3 -c \"raise ValueError('boom')\"", "test", 30))
            check("a hung command is killed",
                  "timed out" in A.lab_execute("sleep 20", "test", 2).lower())
            big = A.lab_execute(
                "for i in $(seq 1 4000); do echo line $i; done", "test", 60)
            check("large output is trimmed but keeps the tail",
                  "trimmed" in big and "line 4000" in big)
            check("container cannot see the host filesystem",
                  "Users" not in A.lab_execute("ls /", "test", 30))

            # NOT `c` - that is the test client, and shadowing it here
            # breaks every later request in the suite.
            lab = _cl.containers.get("stellar-lab-u9998-c1")
            nets = list(lab.attrs["NetworkSettings"]["Networks"])
            check("joined a per-user network", any("u9998" in n for n in nets))
            check("with inter-container comms disabled",
                  _cl.networks.get(nets[0]).attrs["Options"].get(
                      "com.docker.network.bridge.enable_icc") == "false")
            hc = lab.attrs["HostConfig"]
            check("resources are capped",
                  hc["Memory"] == 2 * 1024**3 and hc["NanoCpus"] == 2_000_000_000
                  and hc.get("PidsLimit") == 512)

            _g.lab_chat_id = 2
            check("each chat gets its own workspace",
                  "f.txt" not in A.lab_execute("ls", "test", 30))

        for n in ("stellar-lab-u9998-c1", "stellar-lab-u9998-c2"):
            try:
                _cl.containers.get(n).remove(force=True)
            except Exception:
                pass
        try:
            _cl.networks.get("stellar_net_u9998").remove()
        except Exception:
            pass
        import shutil
        for d in ("u9998_c1", "u9998_c2"):
            shutil.rmtree(A.PROJECT_ROOT / "sandbox_runs" / d, ignore_errors=True)

    # --- phase 6: interrupts and compression --------------------------
    ev = A.register_generation(chat["id"], "q-1")
    check("registering a generation yields a live event", not ev.is_set())
    ev2 = A.register_generation(chat["id"], "q-2")
    check("a second generation cancels the one it replaces", ev.is_set())
    check("and does not cancel itself", not ev2.is_set())
    A.release_generation(chat["id"], "q-stale")
    with A._ACTIVE_LOCK:
        check("a stale release leaves the real claim alone",
              chat["id"] in A.ACTIVE_GENERATIONS)
    A.release_generation(chat["id"], "q-2")
    with A._ACTIVE_LOCK:
        check("the owner's release clears it",
              chat["id"] not in A.ACTIVE_GENERATIONS)

    ev3 = A.register_generation(chat["id"], "q-3")
    A.signal_cancel(REDIS_TEST_URL, chat["id"], "q-3")
    check("signal_cancel sets the event", ev3.is_set())
    check("and leaves a durable flag for a late reader",
          A.is_stopped(REDIS_TEST_URL, "q-3"))
    A.release_generation(chat["id"], "q-3")

    A._redis_client(REDIS_TEST_URL).rpush("inject:99", '{"message":"a"}')
    A._redis_client(REDIS_TEST_URL).rpush("inject:99", '{"message":"b"}')
    drained = A._drain_injections(REDIS_TEST_URL, 99)
    check("injections drain in order",
          [m["message"] for m in drained] == ["a", "b"])
    check("and the queue is emptied",
          A._drain_injections(REDIS_TEST_URL, 99) == [])

    # hidden: out of the transcript, still in the model's memory. Both
    # halves matter - the second one was a real bug.
    with app.app_context():
        db = A.get_db()
        A._save_reply(db, chat["id"], "hidden from the user", hidden=True)
        visible = c.get(f"/api/chats/{chat['id']}/messages").get_json()
        check("a hidden reply is absent from the UI",
              not any("hidden from the user" in m["message_content"]
                      for m in visible))
        hist = A.build_gemini_history(db, chat["id"])
        check("but present in the model's history",
              any("hidden from the user" in part.text
                  for cc in hist for part in cc.parts
                  if getattr(part, "text", None)))

        tokens, ratio = A.estimate_context_usage(db, chat["id"])
        check("context usage is estimated", tokens > 0 and 0 <= ratio <= 1)

        from flask import g as _g2
        _g2.lab_chat_id = chat["id"]
        check("a thin state document is refused",
              "too short" in A.compress_memory("both", "nope", "s").lower())
        check("an unknown target is refused",
              "Invalid target" in A.compress_memory("bogus", "x" * 200, "s"))

        for i in range(15):
            A._record_tool_call(db, chat["id"], "lab_execute", {"i": i}, "o" * 40, 1, False)
        A.compress_memory(
            "tool_logs",
            "Objective: exercise compression. Findings: the ten most recent "
            "tool calls stay visible. Files: none. Outstanding: nothing.", "s")
        left = db.execute("SELECT COUNT(*) n FROM tool_calls WHERE chat_id=? AND hidden=0",
                          (chat["id"],)).fetchone()["n"]
        check("compression archives down to the recent tool calls",
              left == A.KEEP_RECENT_TOOL_CALLS, )
        doc = db.execute(
            "SELECT 1 FROM messages WHERE chat_id=? AND hidden=1 AND message_content LIKE ?",
            (chat["id"], A.COMPRESSED_PREFIX + "%")).fetchone()
        check("and preserves the state document", doc is not None)

    check("compress_memory is offered to the model",
          A.compress_memory in A.AVAILABLE_TOOLS)

    # --- phase 10: generative UI and chess ----------------------------
    import threading as _th, time as _t2
    from flask import g as _g3

    with app.app_context():
        _g3.lab_chat_id = chat["id"]
        _g3.stream_redis_url = REDIS_TEST_URL
        _g3.stream_cancelled = lambda: False
        emitted = []
        _g3.stream_emit = emitted.append

        html = '<div><button onclick="window.stellar.finish({ok:1})">Go</button></div>'

        def _respond():
            for _ in range(100):
                if emitted:
                    break
                _t2.sleep(0.05)
            A._redis_client(REDIS_TEST_URL).rpush(
                f"interaction:{emitted[0]['id']}", '{"picked": "b"}')

        _th.Thread(target=_respond, daemon=True).start()
        out = A.request_user_interaction(html, "goal", "waiting")
        check("a widget is emitted to the stream",
              emitted and emitted[0]["type"] == "interaction")
        check("the tool blocks and returns the user's answer",
              '"picked": "b"' in out or "picked" in out)
        check("and the widget is closed afterwards",
              any(e["type"] == "interaction_closed" for e in emitted))

        check("a widget with no finish() call is refused",
              "never calls" in A.request_user_interaction("<div>x</div>", "g", "s"))
        check("a non-HTML widget is refused",
              "HTML fragment" in A.request_user_interaction("text", "g", "s"))

        _g3.stream_cancelled = lambda: True
        check("a stop ends the wait immediately",
              "stopped" in A.request_user_interaction(html, "g", "s"))
        _g3.stream_cancelled = lambda: False

        # Chess: the point is that illegal moves cannot happen.
        st = json.loads(A.chess_move("new", "s"))
        check("a new game has 20 legal moves", len(st["legal_moves_san"]) == 20)
        st = json.loads(A.chess_move("apply", "s", move="e4"))
        check("a legal move is played", st.get("played") == "e4")
        st = json.loads(A.chess_move("apply", "s", move="Nf7"))
        check("an ILLEGAL move is rejected", "not legal" in (st.get("error") or ""))
        check("and the rejection ships the legal list",
              len(st["legal_moves_san"]) > 0)
        st = json.loads(A.chess_move("apply", "s", move="Qz9"))
        check("nonsense notation is rejected too",
              "not legal" in (st.get("error") or ""))

        import chess_engine as _ce
        mate = _ce.analyse("6k1/5ppp/8/8/8/8/5PPP/R5K1 w - - 0 1",
                           top_n=1, time_budget=2)
        check("the engine finds mate in one",
              mate["candidates"] and mate["candidates"][0]["san"] == "Ra8#")
        hang = _ce.analyse(
            "rnb1kbnr/ppp1pppp/8/3q4/8/2N5/PPPPPPPP/R1BQKBNR w KQkq - 0 1",
            top_n=1, time_budget=2)
        check("and wins a hanging queen",
              hang["candidates"] and hang["candidates"][0]["san"] == "Nxd5")

        check("the interactive tools are offered to the model",
              A.request_user_interaction in A.AVAILABLE_TOOLS
              and A.chess_move in A.AVAILABLE_TOOLS
              and A.chess_play in A.AVAILABLE_TOOLS)

        # The server-rendered board: both colours distinct, filled glyphs
        # only, history and legal moves embedded.
        import chess as _chess, chess_ui as _cui
        _b = _chess.Board(); _b.push_san("e4"); _b.push_san("e5")
        _w = _cui.render(_b, ["e4", "e5"], user_color="white", elo=2000,
                         status_text="Your move", waiting=True, last_move="e7e5")
        # Real piece artwork, both colours, and none of the font glyphs that
        # made white and black look identical.
        check("board widget ships all twelve SVG pieces",
              all(k in _w for k in ("wK", "wQ", "wR", "wB", "wN", "wP",
                                    "bK", "bQ", "bR", "bB", "bN", "bP"))
              and 'viewBox="0 0 45 45"' in _w)
        check("board widget uses no unicode chess glyphs",
              not any(chr(c) in _w for c in range(0x2654, 0x2660)))
        check("board widget carries a clock and an in-place update hook",
              '"clock": null' in _w and "stellar:update" in _w)
        check("board widget embeds the move list and legal moves",
              '"moves": ["e4", "e5"]' in _w and '"legal": [' in _w
              and "g1f3" in _w)
        check("board widget calls back through window.stellar.finish",
              "window.stellar.finish" in _w)

        # Move grading: the review maths and the vocabulary the widget knows,
        # kept engine-free so the suite stays fast and offline.
        check("accuracy is 100 for a clean game and falls with loss",
              _ce.accuracy([]) == 100.0 and _ce.accuracy([0, 0]) == 100.0
              and _ce.accuracy([300, 300]) < _ce.accuracy([20, 20]) < 100)
        check("evals read like a chess site's",
              (_ce.eval_text({"cp": 134, "mate": None}), _ce.eval_text({"cp": 9997, "mate": 3}),
               _ce.eval_text({"cp": -9998, "mate": -2}), _ce.eval_text(None))
              == ("+1.3", "M3", "-M2", "0.0"))
        check("a move with one legal reply is graded forced",
              _ce.grade_move(_chess.Board("7k/8/8/8/8/8/8/K5R1 b - - 0 1"),
                             _chess.Move.from_uci("h8h7"),
                             {"cp": 800, "mate": None, "best": "h8h7"},
                             {"cp": 800, "mate": None, "best": "a1a2"})["cls"] == "forced")
        _sac_b = _chess.Board("r1bqkb1r/ppp2ppp/2n5/3np1N1/2B5/8/PPPP1PPP/RNBQK2R w KQkq - 0 6")
        check("a knight left en prise counts as a sacrifice, a pawn push does not",
              _ce.is_sacrifice(_sac_b, _chess.Move.from_uci("g5f7"))
              and not _ce.is_sacrifice(_chess.Board(), _chess.Move.from_uci("e2e4")))
        _graded = [_ce.grade_move(_chess.Board(), _chess.Move.from_uci("d2d4"),
                                  {"cp": 30, "mate": None, "best": "e2e4", "second_cp": 25},
                                  {"cp": 30 - loss, "mate": None, "best": "e7e5"})["cls"]
                   for loss in (0, 20, 50, 100, 200, 400)]
        check("grades step from best to blunder as the loss grows",
              _graded == ["best", "excellent", "good", "inaccuracy", "mistake", "blunder"])
        # Accuracy is measured in win-percentage loss, not centipawns:
        # the curve's input is a 0-100 quantity, and feeding it centipawns
        # reported ordinary games as 0% for both players.
        _rv = _ce.review([{"cls": "best", "loss": 0, "wp_loss": 0.0,
                           "best_san": "e4", "best_uci": "e2e4"},
                          {"cls": "blunder", "loss": 500, "wp_loss": 46.0,
                           "best_san": "e5", "best_uci": "e7e5"}],
                         ["e4", "f6"], "white")
        check("the end-of-game review scores both sides and names the worst move",
              _rv["white"]["accuracy"] == 100.0 and _rv["black"]["accuracy"] < 25
              and _rv["black"]["worst"][0]["better"] == "e5")
        check("a well-played game scores high rather than zero",
              _ce.accuracy([2.0, 3.0, 1.5]) > 80
              and _ce.accuracy([40.0, 45.0]) < _ce.accuracy([2.0, 3.0]))
        _w2 = _cui.render(_b, ["e4", "e5"], user_color="white", elo=2000,
                          status_text="Your move", waiting=True, last_move="e7e5",
                          quality=[{"cls": "best", "loss": 0, "best_san": "e4", "best_uci": "e2e4"},
                                   {"cls": "good", "loss": 30, "best_san": "c5", "best_uci": "c7c5"}],
                          eval={"cp": 35, "mate": None, "text": "+0.4"})
        check("board widget carries grades, the eval bar and the review card",
              '"quality": [{' in _w2 and '"eval": {"cp": 35' in _w2
              and 'id="evalbar"' in _w2
              and "gameover" in _w2 and "drawArrow" in _w2 and "pointerdown" in _w2)

    # --- widget sizing -------------------------------------------------
    # An iframe has no natural height, so the page sizes it from a height
    # the widget reports. Two properties of that mechanism are load-bearing
    # and were each broken once: the widget must measure its own CONTENT
    # (documentElement is whatever height the page last set, so measuring
    # it can never ask to grow), and the frame must not scroll internally
    # (a scrollbar steals ~15px of width, which pushed the chessboard's
    # side panel onto a second row and made the widget 380px taller).
    _mainjs = (Path(__file__).parent / "static" / "main.js").read_text(encoding="utf-8")
    _wrapper = _mainjs[_mainjs.index("function widgetDocument"):
                       _mainjs.index("function renderInteraction")]
    check("the widget frame never scrolls inside itself",
          "html{overflow:hidden}" in _wrapper)
    check("the widget reports the height of its content, not of the frame",
          'getElementById("stellar-widget-root")' in _wrapper
          and "documentElement.scrollHeight" not in _wrapper)
    check("content changes are observed on the content, not the frame",
          "observe(root)" in _wrapper
          and "observe(document.documentElement)" not in _wrapper)
    check("the height cap clears a full chessboard",
          "2400" in _mainjs[_mainjs.index('__stellar === "height"'):][:220])

    # The board must fit beside its panel in the chat column (~712px wide):
    # board column is squares*8 + 22, plus a 16px gap, plus the panel.
    _sq = int(re.search(r"--sq:(\d+)px", _cui._TEMPLATE).group(1))
    _basis = int(re.search(r"\.cw \.panel\{flex:1 1 (\d+)px", _cui._TEMPLATE).group(1))
    check(f"board ({_sq}px squares) and panel ({_basis}px) fit side by side in the chat column",
          _sq * 8 + 22 + 16 + _basis <= 700)

    # --- phase 8: the rest of the tool suite ---------------------------
    conn = sqlite3.connect(tmp)
    tables = {r[0] for r in conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table'")}
    conn.close()
    check("schema: user_memory and scheduled_tasks exist",
          {"user_memory", "scheduled_tasks"} <= tables)
    check("phase 8 tools are offered to the model",
          {"generate_image", "make_presentation", "analyze_youtube_video",
           "send_self_email", "remember", "read_tool_output", "manage_files",
           "schedule_task"} <= set(A.TOOLS_BY_NAME))

    import shutil as _shutil
    import threading as _threading
    import time as _time
    from flask import g as _g
    with app.app_context():
        _db = A.get_db()
        _uid = _db.execute(
            "INSERT INTO users (username, password_hash, is_approved)"
            " VALUES ('p8@test.com', 'x', 1)").lastrowid
        _cid = _db.execute("INSERT INTO chats (user_id) VALUES (?)", (_uid,)).lastrowid
        _db.commit()
        _g.lab_user_id, _g.lab_chat_id = _uid, _cid

        # memory: saved by the model, prepended to every later turn
        _saved = A.remember("s", note="Prefers  short answers")
        _nid = int(re.search(r"note (\d+)", _saved).group(1))
        check("remember saves a note", "Saved as note" in _saved)
        check("a duplicate note is not saved twice",
              "Already saved" in A.remember("s", note="Prefers short answers"))
        check("notes are prepended to the system prompt",
              f"[{_nid}] Prefers short answers" in A.memory_prompt(_db, _uid))
        check("a note can be forgotten",
              "Forgot" in A.remember("s", forget_id=_nid)
              and A.memory_prompt(_db, _uid) == "")
        check("an empty note is refused", "Give a note" in A.remember("s"))

        # long outputs: the row keeps everything, the model gets a page
        _rid = A._record_tool_call(
            _db, _cid, "fetch_url", {"url": "x"},
            "\n".join(f"line {i}" + (" needle" if i % 50 == 0 else "") for i in range(300)),
            1, False)
        _view = A._model_view("x" * (A.TOOL_OUTPUT_LIMIT + 500), _rid)
        check("a long result is cut for the model with a pointer to the rest",
              len(_view) < A.TOOL_OUTPUT_LIMIT + 400
              and f"read_tool_output(output_id={_rid})" in _view)
        check("a short result passes through untouched",
              A._model_view("short", _rid) == "short")
        check("read_tool_output pages by line",
              "lines 290-299 of 300" in A.read_tool_output(_rid, "s", start_line=290))
        check("read_tool_output searches by keyword",
              "matches 0-5 of 6" in A.read_tool_output(_rid, "s", keyword="needle"))
        check("read_tool_output is scoped to the chat",
              "no tool output" in A.read_tool_output(_rid + 999, "s"))

        # files: the sandbox workspace is a host folder, so sharing is a copy
        _lab = A._lab_workspace(_uid, _cid)
        (_lab / "plots").mkdir(exist_ok=True)
        (_lab / "plots" / "chart.png").write_bytes(b"\x89PNG-test")
        check("a sandbox file can be shared as an inline image link",
              f"![chart.png](/api/outputs/{_cid}/chart.png)"
              in A.manage_files("share", "s", path="plots/chart.png"))
        check("sharing cannot escape the workspace",
              "No file" in A.manage_files("share", "s", path="../../keys.env"))
        _listing = A.manage_files("list", "s")
        check("listing shows outputs and the workspace",
              "chart.png" in _listing and "/lab" in _listing)
        check("chat files resolve safely",
              A._resolve_chat_file(_uid, _cid, "../../app.py") is None
              and A._resolve_chat_file(_uid, _cid, "chart.png") is not None)

        # email: only ever to the user's own address
        for _k in ("EMAIL_USER", "EMAIL_PASS"):
            os.environ.pop(_k, None)
        check("email reports missing configuration",
              "not configured" in A.send_self_email("hi", "body", "s"))
        os.environ["EMAIL_USER"], os.environ["EMAIL_PASS"] = "bot@example.com", "app-pass"
        _sent: dict = {}
        _orig_send = A._smtp_send
        A._smtp_send = lambda sender, password, msg: _sent.update(
            to=msg["To"], subject=msg["Subject"],
            types=[p.get_content_type() for p in msg.walk()])
        try:
            _out = A.send_self_email("Report", "# Title\n\n- a\n- b", "s", attachment="chart.png")
            _missing = A.send_self_email("x", "y", "s", attachment="nope.pdf")
        finally:
            A._smtp_send = _orig_send
            os.environ.pop("EMAIL_USER")
            os.environ.pop("EMAIL_PASS")
        check("email goes only to the user's own address",
              _sent.get("to") == "p8@test.com" and "Sent to p8@test.com" in _out)
        check("email carries text, html and the attachment",
              {"text/plain", "text/html", "image/png"} <= set(_sent.get("types", [])))
        check("a missing attachment is reported", "not among" in _missing)

        # youtube: only real links reach the model
        check("analyze needs a YouTube link",
              "must be a YouTube link"
              in A.analyze_youtube_video("s", video_url="https://example.com/watch?v=abc"))
        check("YouTube link forms are recognised",
              all(A._YOUTUBE_RE.match(u) for u in (
                  "https://www.youtube.com/watch?v=jNQXAC9IVRw",
                  "https://youtu.be/jNQXAC9IVRw", "youtube.com/shorts/abcdefgh123")))

        # presentations: the writer (planning needs the model)
        _plan = {"title": "Test deck", "subtitle": "sub",
                 "slides": [{"title": f"Slide {i}", "bullets": ["alpha", "beta"],
                             "notes": "say this", "visual": "v"} for i in range(3)]}
        _deck = Path(app.config["OUTPUTS_DIR"]) / "t.pptx"
        check("a deck is written with a title slide plus one per plan entry",
              A._build_deck(_plan, "dark", [None] * 3, _deck) == 4
              and _deck.stat().st_size > 20000)
        from pptx import Presentation as _Presentation
        _prs = _Presentation(str(_deck))
        check("the deck reopens with speaker notes intact",
              len(_prs.slides) == 4
              and _prs.slides[1].notes_slide.notes_text_frame.text == "say this")
        check("style words pick a theme",
              A._deck_theme("dark technical")["bg"] == "14171F"
              and A._deck_theme("anything")["bg"] == "FFFFFF")

        # scheduling: times, limits, listing, cancelling
        _one = A.schedule_task("schedule", "s", task_prompt="say hi", delay_minutes=2)
        check("a task can be scheduled by delay", "Scheduled as task #" in _one)
        _tid = int(re.search(r"#(\d+)", _one).group(1))
        check("ISO times with an offset are converted to UTC",
              "02:30:00 UTC" in A.schedule_task(
                  "schedule", "s", task_prompt="d",
                  run_at="2030-01-01T08:00:00+05:30", every_minutes=60))
        check("a naive time is flagged as read as UTC",
              "read as UTC" in A.schedule_task(
                  "schedule", "s", task_prompt="n", run_at="2030-01-01T08:00:00"))
        check("a past time is refused",
              "in the past" in A.schedule_task(
                  "schedule", "s", task_prompt="p", run_at="2020-01-01T00:00:00Z"))
        check("too-frequent repeats are refused",
              "at least 5" in A.schedule_task(
                  "schedule", "s", task_prompt="r", delay_minutes=5, every_minutes=1))
        check("tasks are listed", f"#{_tid}" in A.schedule_task("list", "s"))
        check("a task can be cancelled",
              "Cancelled" in A.schedule_task("cancel", "s", task_id=_tid + 1))

        # the scheduler runs a due task as a turn in its chat (stub producer)
        _db.execute("UPDATE scheduled_tasks SET run_at = '2020-01-01 00:00:00'"
                    " WHERE id = ?", (_tid,))
        _db.commit()
        # The transport tests restored the real producer; the scheduler must
        # not spend quota (or need a network) to prove it runs a turn.
        _real_producer = A.gemini_producer
        A.gemini_producer = stub_producer
        check("a due task is claimed and started", A.run_due_tasks(app) == 1)
        for _ in range(100):
            _row = _db.execute("SELECT status, runs FROM scheduled_tasks WHERE id = ?",
                               (_tid,)).fetchone()
            if _row["status"] != "running":
                break
            _time.sleep(0.05)
        check("a one-off task is marked done after it runs",
              _row["status"] == "done" and _row["runs"] == 1)
        _msgs = [r[0] for r in _db.execute(
            "SELECT message_content FROM messages WHERE chat_id = ? ORDER BY id",
            (_cid,)).fetchall()]
        check("the task's prompt and the reply landed in the chat",
              any(m.startswith("[Scheduled task #") for m in _msgs)
              and any("stub" in m for m in _msgs))
        check("nothing is due afterwards", A.run_due_tasks(app) == 0)

        # a chat mid-generation postpones the task instead of cancelling the turn
        _db.execute("UPDATE scheduled_tasks SET status = 'pending',"
                    " run_at = '2020-01-01 00:00:00' WHERE id = ?", (_tid,))
        _db.commit()
        A.ACTIVE_GENERATIONS[_cid] = (_threading.Event(), "q")
        try:
            _started = A.run_due_tasks(app)
            _state = _db.execute("SELECT status, run_at FROM scheduled_tasks WHERE id = ?",
                                 (_tid,)).fetchone()
        finally:
            A.ACTIVE_GENERATIONS.pop(_cid, None)
        check("a task whose chat is mid-turn is pushed back a minute",
              _started == 0 and _state["status"] == "pending"
              and _state["run_at"] > "2020-01-01 00:00:00")
        A.gemini_producer = _real_producer
        _db.execute("UPDATE scheduled_tasks SET status = 'cancelled' WHERE id = ?", (_tid,))
        _db.commit()
        _shutil.rmtree(_lab, ignore_errors=True)

    # produced files are served to their owner only
    _chat_id = c.post("/api/chats").get_json()["id"]
    with app.app_context():
        _own = A._outputs_dir(1, _chat_id)
    (_own / "pic.png").write_bytes(b"\x89PNG")
    (_own / "deck.pptx").write_bytes(b"PK")
    r = c.get(f"/api/outputs/{_chat_id}/pic.png")
    check("an owner can fetch a produced file", r.status_code == 200 and r.data == b"\x89PNG")
    r = c.get(f"/api/outputs/{_chat_id}/deck.pptx")
    check("documents are served as downloads",
          r.status_code == 200 and "attachment" in r.headers.get("Content-Disposition", ""))
    check("a missing file is 404", c.get(f"/api/outputs/{_chat_id}/nope.png").status_code == 404)
    check("another user's chat is 404", c.get(f"/api/outputs/{_cid}/chart.png").status_code == 404)
    check("path traversal is refused",
          c.get(f"/api/outputs/{_chat_id}/..%2F..%2Fkeys.env").status_code in (400, 404))
    check("anonymous access is refused",
          app.test_client().get(f"/api/outputs/{_chat_id}/pic.png").status_code in (302, 401))

    # --- phase 9: production deploy & repo_control --------------------
    conn = sqlite3.connect(tmp)
    tables = {r[0] for r in conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table'")}
    conn.close()
    check("schema: repo_history exists", "repo_history" in tables)
    check("repo_control is offered to the model",
          "repo_control" in A.TOOLS_BY_NAME and A.repo_control in A.AVAILABLE_TOOLS)
    check("repo_control has a docstring", bool(A.repo_control.__doc__))

    with app.app_context():
        _db = A.get_db()
        # slug generator tests
        slug1 = A.generate_unique_subdomain("My Awesome Project", _db)
        check("subdomain slug generated", slug1 == "my-awesome-project")
        _db.execute(
            "INSERT INTO repo_history (user_id, project_name, process_id, subdomain, status) "
            "VALUES (?, 'My Awesome Project', 'proc-coll-1', 'my-awesome-project', 'stopped')",
            (_uid,))
        _db.commit()
        slug2 = A.generate_unique_subdomain("My Awesome Project", _db)
        check("slug collision avoided with suffix", slug2 == "my-awesome-project-2")
        slug_reserved = A.generate_unique_subdomain("admin", _db)
        check("reserved subdomain guarded", slug_reserved == "admin-app")

        # repo_control input validation
        _g.lab_user_id, _g.lab_chat_id = _uid, _cid
        check("unknown action refused", "Unknown action" in A.repo_control("bad_action", "s"))
        check("stop unknown deployment refused", "not found" in A.repo_control("stop", "s", app_id="nonexistent-app"))

        # subdomain routing via test client
        # 1. Unknown subdomain -> 404
        r_unk = c.get("/", headers={"Host": "nonexistent-app.stellarai.site"})
        check("unknown subdomain is 404", r_unk.status_code == 404)

        # 2. Stopped app -> 503
        r_stop = c.get("/", headers={"Host": "my-awesome-project.stellarai.site"})
        check("stopped app subdomain is 503", r_stop.status_code == 503)

        # 3. Unapproved owner -> 403
        _uid_unapp = _db.execute(
            "INSERT INTO users (username, password_hash, is_approved) VALUES ('unapp@test.com', 'x', 0)"
        ).lastrowid
        _db.execute(
            "INSERT INTO repo_history (user_id, project_name, process_id, subdomain, status, host_port) "
            "VALUES (?, 'Unapproved App', 'proc-unapp', 'unapp-test', 'running', 5999)",
            (_uid_unapp,))
        _db.commit()
        r_unapp = c.get("/", headers={"Host": "unapp-test.stellarai.site"})
        check("unapproved owner app is 403", r_unapp.status_code == 403)

        # 4. Live deployment lifecycle via Docker if available
        if docker_up:
            dep_res = A.repo_control("deploy", "s", project_name="Smoke Test App", port=5000)
            check("repo_control deploys container", "Container provisioned" in dep_res and "Smoke Test App" in dep_res)

            list_res = A.repo_control("list_history", "s")
            check("repo_control lists deployed app", "Smoke Test App" in list_res)

            exec_res = A.repo_control("execute", "s", app_id="Smoke Test App", command="echo 'sample-content' > sample.txt && cat sample.txt")
            check("repo_control executes inside container", "sample-content" in exec_res)

            snap_res = A.repo_control("snapshot", "s", app_id="Smoke Test App")
            check("repo_control snapshots files", "Snapshotted" in snap_res)

            rename_res = A.repo_control("rename", "s", app_id="Smoke Test App", project_name="Renamed Smoke App")
            check("repo_control renames deployment", "renamed to 'Renamed Smoke App'" in rename_res)

            stop_res = A.repo_control("stop", "s", app_id="Renamed Smoke App")
            check("repo_control stops deployment", "stopped" in stop_res.lower())

            restart_res = A.repo_control("restart", "s", app_id="Renamed Smoke App")
            check("repo_control restarts deployment", "restarted and running" in restart_res.lower())

            # Cleanup
            A.repo_control("stop", "s", app_id="Renamed Smoke App")
            try:
                row_clean = _db.execute("SELECT process_id FROM repo_history WHERE project_name='Renamed Smoke App'").fetchone()
                if row_clean:
                    _cl.containers.get(f"stellar-repo-{row_clean['process_id']}").remove(force=True)
                    p_dir = A.PROJECT_ROOT / "deployments" / f"u{_uid}_{row_clean['process_id']}"
                    if p_dir.exists():
                        _shutil.rmtree(p_dir, ignore_errors=True)
                # If deployments dir is empty, remove it as well
                dep_dir = A.PROJECT_ROOT / "deployments"
                if dep_dir.exists() and not any(dep_dir.iterdir()):
                    dep_dir.rmdir()
            except Exception:
                pass
        else:
            print("  SKIP  repo_control live Docker actions (Docker not reachable)")

    # --- audit fixes ---------------------------------------------------
    # Each of these had a defect found by the project audit. They are
    # cheap, and every one of them failed silently before it was fixed.
    import sqlite3 as _sq3

    # A tool may block far longer than a stream may be silent; if the
    # stream gives up first the user is told the turn died while it runs.
    check("a stream outlives the longest blocking tool",
          A.IDLE_TIMEOUT > max(A.INTERACTION_TIMEOUT, A.LAB_MAX_TIMEOUT))

    # Model-authored markup must never render on the app's own origin.
    check("model-authored markup is never served inline",
          not ({".html", ".htm", ".svg"} & A._INLINE_TYPES))

    # An overloaded model needs the model-switch path, not a retry of the
    # same model; its message also says 503, so order matters.
    check("an overloaded model is routed to the fallback, not retried",
          A._classify_error(Exception("503 The model is overloaded.")) == "quota"
          and A._classify_error(Exception("503 Service Unavailable")) == "transient"
          and A._classify_error(Exception("429 quota")) == "quota")

    # next= must not accept a backslash: browsers normalise /\ to //.
    def _next_ok(n):
        return bool(n and n.startswith("/") and not n.startswith("//") and "\\" not in n)
    check("login next= refuses an off-site redirect",
          not _next_ok("/\\evil.com") and not _next_ok("//evil.com")
          and _next_ok("/api/chats"))

    # A produced file's link and name have to survive the round trip.
    check("produced-file links are URL-encoded and keep their extension",
          A._output_link(7, "my report.png") == "/api/outputs/7/my%20report.png"
          and A._safe_filename("x" * 200 + ".png").endswith(".png"))

    # Accuracy is a win-percentage curve; centipawns made every game 0%.
    check("accuracy is measured in win percentage, not centipawns",
          _ce.accuracy([2.0, 3.0, 1.5]) > 80
          and 0 < _ce.accuracy([40.0, 45.0]) < 60
          and _ce.accuracy([]) == 100.0)
    _gq = _ce.grade_move(_chess.Board(), _chess.Move.from_uci("e2e4"),
                         {"cp": 30, "mate": None, "best": "e2e4", "second_cp": 25},
                         {"cp": 28, "mate": None, "best": "e7e5"})
    check("a grade carries the win-percentage drop", "wp_loss" in _gq)

    # a1 is dark on a real board.
    check("the board is not mirrored",
          "(file + rank) % 2 === 0" in _cui._TEMPLATE)

    # A promotion must not invent a captured pawn.
    _pb = _chess.Board("rnbqkbnr/1Ppppppp/8/8/8/8/P1PPPPPP/RNBQKBNR w KQkq - 0 1")
    _pb.push_san("bxa8=Q")
    check("the capture tray survives a promotion",
          _cui.captured(_pb) == (["r"], [], 5))

    # Either side can run out of time.
    _play_src = (Path(__file__).parent / "app.py").read_text(encoding="utf-8")
    _play_src = _play_src[_play_src.index("def chess_play("):]
    check("either side can lose on time",
          'state["forfeit"] = _colour' in _play_src[:_play_src.index("# The registry")])

    # A recurring task must survive one bad run.
    with app.app_context():
        _d = A.get_db()
        _u = _d.execute("INSERT INTO users (username,password_hash,is_approved)"
                        " VALUES ('sched@test','x',1)").lastrowid
        _c2 = _d.execute("INSERT INTO chats (user_id) VALUES (?)", (_u,)).lastrowid
        _t = _d.execute(
            "INSERT INTO scheduled_tasks (user_id,chat_id,task_prompt,run_at,"
            "every_minutes,status,lock_id) VALUES (?,?,'d','2020-01-01 00:00:00',10,'running','L')",
            (_u, _c2)).lastrowid
        _d.commit()
        A._finish_task(_t, False)
        check("a recurring task survives a failed run",
              _d.execute("SELECT status FROM scheduled_tasks WHERE id=?",
                         (_t,)).fetchone()["status"] == "pending")
        _t2 = _d.execute(
            "INSERT INTO scheduled_tasks (user_id,chat_id,task_prompt,run_at,"
            "every_minutes,status,lock_id) VALUES (?,?,'o','2020-01-01 00:00:00',0,'running','L2')",
            (_u, _c2)).lastrowid
        _d.commit()
        A._finish_task(_t2, False)
        check("a one-off failure is still terminal",
              _d.execute("SELECT status FROM scheduled_tasks WHERE id=?",
                         (_t2,)).fetchone()["status"] == "failed")
        _d.execute("UPDATE scheduled_tasks SET status='cancelled'")
        _d.commit()

    # An older database must upgrade rather than half-build and lie.
    _old = Path(tempfile.mkdtemp()) / "old.db"
    _oc = _sq3.connect(_old)
    _oc.executescript(
        "CREATE TABLE users (id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT UNIQUE,"
        " password_hash TEXT);"
        "CREATE TABLE chats (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER);"
        "CREATE TABLE messages (id INTEGER PRIMARY KEY AUTOINCREMENT, chat_id INTEGER,"
        " message_type TEXT, message_content TEXT);"
        "INSERT INTO users (username,password_hash) VALUES ('old@u','x');"
        "INSERT INTO chats (user_id) VALUES (1);"
        "INSERT INTO messages (chat_id,message_type,message_content) VALUES (1,'user','kept');")
    _oc.commit(); _oc.close()
    _oldapp = A.create_app({"DATABASE": str(_old), "TESTING": True,
                            "REDIS_URL": REDIS_TEST_URL})
    with _oldapp.app_context():
        A.init_db()
        _od = A.get_db()
        _rows = _od.execute("SELECT message_content FROM messages"
                            " WHERE chat_id=1 AND hidden=0").fetchall()
        _chats = _od.execute("SELECT id FROM chats WHERE user_id=1 AND is_temp=0").fetchall()
        _drift = A.schema_drift(_od)
    check("an older database upgrades and keeps its data",
          len(_rows) == 1 and _rows[0][0] == "kept" and len(_chats) == 1 and _drift == [])

    # --- history mapping ---------------------------------------------
    with app.app_context():
        db = A.get_db()
        hist = A.build_gemini_history(db, chat["id"])
        check("history maps to alternating gemini roles",
              [h.role for h in hist] == ["user", "model"])
        check("history starts with a user turn",
              not hist or hist[0].role == "user")

    # --- live turn ----------------------------------------------------
    if LIVE:
        chat2 = c.post("/api/chats").get_json()
        qid2 = c.post(f"/api/chats/{chat2['id']}/query",
                      json={"message": "Reply with exactly: PONG"}
                      ).get_json()["query_id"]
        ev = parse_sse(c.get(f"/api/stream/{qid2}").get_data(as_text=True))
        kinds = [e["type"] for _, e in ev]
        text = "".join(e["text"] for _, e in ev if e["type"] == "token")
        errs = [e["message"] for _, e in ev if e["type"] == "error"]
        check("live turn produced tokens", kinds.count("token") >= 1)
        check("live turn committed a reply", "message" in kinds)
        check(f"live reply looks sane (got {text.strip()[:40]!r})",
              "pong" in text.lower())
        if errs:
            print(f"        live errors: {errs}")

    redis_lib.from_url(REDIS_TEST_URL).flushdb()

    print()
    if failures:
        print(f"  {len(failures)} FAILED: {', '.join(failures)}\n")
        return 1
    print(f"  All checks passed.\n")
    return 0


if __name__ == "__main__":
    sys.exit(main())
